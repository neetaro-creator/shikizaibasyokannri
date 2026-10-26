"""点検表(xlsx)から資器材 CSV の下書きを作る。

使い方: python3 tools/import_xlsx.py <xlsxのあるディレクトリ> <出力ディレクトリ> [--existing <今のdata>]
  --existing を付けると、同じ品名・場所名には今と同じ番号を使い、手入力した欄を引き継ぐ(点検表の改訂時はこれを使う)。

帳票は「品名列+数量列」の列帯が横に並ぶ。列帯ごとに上から読み、
数量のない見出しで場所を更新し、品名+数量を在庫行にする。
判断に自信がない行は review.csv に出す。
"""
import csv
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
import pykakasi
from openpyxl.utils import column_index_from_string

CONFIG = json.loads((Path(__file__).parent / "import_config.json").read_text(encoding="utf-8"))
HEADING = re.compile(CONFIG["heading_pattern"])
SUBLEVEL = re.compile(CONFIG["sublevel_pattern"])
INSPECTION = re.compile(CONFIG["inspection_pattern"])
NOT_ITEMS = [unicodedata.normalize("NFKC", n) for n in CONFIG["not_items"]]
HEADER_ROW = re.compile(r"積載場所")
KANJI = re.compile(r"[\u4e00-\u9fff]")
KAKASI = pykakasi.kakasi()


def reading(name):
    """漢字を含む品名のひらがな読み(検索の通称用)。漢字がなければ空。"""
    if not KANJI.search(name):
        return ""
    return "".join(t["hira"] for t in KAKASI.convert(name)).replace(" ", "")


def clean(s):
    s = unicodedata.normalize("NFKC", str(s))
    s = s.split("※")[0]
    return re.sub(r"\s+", " ", s).strip()


def loc_clean(s):
    return re.sub(r"\s+", "", clean(s))


def key(s):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", s)).lower()


def col_range(spec):
    a, b = spec.split(":")
    return list(range(column_index_from_string(a), column_index_from_string(b) + 1))


def parse_qty(cells):
    """数量セル群 → (qty, note, inspection)。"""
    nums, notes, insp = [], [], []
    for c in cells:
        t = re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(c)))
        if not t:
            continue
        if INSPECTION.search(t):
            insp.append(t)
            continue
        m = re.fullmatch(r"計?(\d+)(本|個|枚|組|双|台|缶|巻)?", t)
        if m:
            nums.append(int(m.group(1)))
        else:
            notes.append(t)
    qty = str(sum(nums)) if nums else ""
    note = " ".join(notes)
    if insp and not nums and not notes:
        note = f"あり(点検:{'・'.join(insp)})"
    return qty, note, insp


def vertical_runs(ws, cols, rows):
    """縦書きラベル(1〜2文字のセルが縦に3つ以上並ぶ)を検出し、行→ラベルを返す。"""
    labels = {}  # (row) -> list[(col, label)]
    for c in cols:
        run = []
        def flush():
            if len(run) >= 3:
                text = "".join(clean(v) for _, v in run).replace("❘", "ー").replace("|", "ー").replace("・", "")
                first, last = run[0][0], run[-1][0]
                for r in range(first, last + 1):
                    labels.setdefault(r, []).append((c, text))
            run.clear()
        for r in rows:
            v = ws.cell(r, c).value
            t = clean(v) if v is not None else ""
            if t and len(t) <= 2 and not SUBLEVEL.match(t):
                if run and r - run[-1][0] > 2:
                    flush()
                run.append((r, t))
            elif t:
                flush()
        flush()
    return labels


class Builder:
    def __init__(self):
        self.vehicles, self.locations, self.items, self.stock, self.review = [], {}, {}, {}, []

    def location(self, vid, name):
        k = (vid, name)
        if k not in self.locations:
            n = sum(1 for v, _ in self.locations if v == vid) + 1
            self.locations[k] = f"{vid}-l{n:02d}"
        return self.locations[k]

    def item(self, name):
        k = key(name)
        if k not in self.items:
            self.items[k] = (f"i{len(self.items) + 1:04d}", name)
        return self.items[k][0]

    def add(self, vid, loc, name, qty, note):
        lid, iid = self.location(vid, loc), self.item(name)
        if (iid, lid) in self.stock:
            q0, n0 = self.stock[(iid, lid)]
            qty = str(int(q0 or 0) + int(qty or 0)) if (q0 or qty) else ""
            note = " ".join(x for x in (n0, note) if x)
        self.stock[(iid, lid)] = (qty, note)


def read_band_by_rows(b, v, ws, band, tcols, qcols):
    """場所の範囲を行で明示した列帯を読む(縦書きラベルが品目の上下中央に書かれた様式用)。

    band["locations"] = [[開始行, 終了行, 場所名], ...]
    band["merge"] = [[開始行, 終了行], ...]  … 1品目が複数行に折り返している行をまとめる
    """
    vid = v["id"]
    merged = {a: z for a, z in band.get("merge", [])}
    inside = {r for a, z in band.get("merge", []) for r in range(a + 1, z + 1)}
    for r0, r1, loc in band["locations"]:
        loc = loc_clean(loc) if " " not in loc else " ".join(loc_clean(x) for x in loc.split(" "))
        for r in range(r0, r1 + 1):
            if r in inside:
                continue
            group = range(r, merged.get(r, r) + 1)
            texts = []
            for c in tcols:
                raw = [ws.cell(g, c).value for g in group if ws.cell(g, c).value is not None]
                t = clean("\n".join(str(x) for x in raw)) if raw else ""
                if t:
                    texts.append((c, t))
            qcells = [ws.cell(g, c).value for g in group for c in qcols if ws.cell(g, c).value is not None]
            if not texts:
                continue
            qty, note, insp = parse_qty(qcells)
            has_qty = bool(qty or note or insp)
            name = texts[-1][1]  # 縦書きラベルの文字は品名より左にある
            if not has_qty and HEADING.search(loc_clean(name)):
                continue  # 見出し行
            if len(name) <= 2 and not qty:
                continue  # 縦書きラベルの文字だけの行(「斧」のように数のある品目は残す)
            cell = f"R{r}"
            if any(key(name).startswith(key(n)) for n in NOT_ITEMS):
                b.review.append([vid, v["sheet"], cell, name, "not_item"])
                continue
            if not has_qty:
                note = "数量記載なし"
                b.review.append([vid, v["sheet"], cell, name, "noqty"])
            if insp:
                b.review.append([vid, v["sheet"], cell, f"{name} [{' '.join(insp)}]", "inspection"])
            b.add(vid, loc, name, qty, note)


def read_vehicle(b, v, src):
    wb = openpyxl.load_workbook(src / v["file"], data_only=True)
    ws = wb[v["sheet"]]
    vid = v["id"]
    b.vehicles.append(v)
    for band in v["bands"]:
        tcols, qcols = col_range(band["text"]), col_range(band["qty"])
        rows = [r for a, z in band["rows"] for r in range(a, z + 1)]
        if "locations" in band:
            read_band_by_rows(b, v, ws, band, tcols, qcols)
            continue
        vlabels = vertical_runs(ws, tcols, rows)
        vcells = {(r, c) for r, lst in vlabels.items() for c, _ in lst}
        major = loc_clean(band["start"]) if band.get("start") else ""
        if major:
            b.review.append([vid, v["sheet"], f"{band['text']}@{rows[0]}", major, "start_inferred"])
        sub, sub_col, container = "", None, ""
        for r in rows:
            texts = [(c, ws.cell(r, c).value) for c in tcols]
            texts = [(c, clean(t)) for c, t in texts if t is not None and (r, c) not in vcells and clean(t)]
            qcells = [ws.cell(r, c).value for c in qcols if ws.cell(r, c).value is not None]
            runs = tuple(t for _, t in sorted(vlabels.get(r, [])))
            if runs:
                container = " ".join(runs)
            if not texts:
                continue
            if HEADER_ROW.search(texts[0][1]):
                continue
            qty, note, insp = parse_qty(qcells)
            has_qty = bool(qty or note or insp)
            if not has_qty and len(texts) == 1:
                c, t = texts[0]
                if SUBLEVEL.match(loc_clean(t)):
                    sub, sub_col = loc_clean(t), c
                    continue
                if HEADING.search(loc_clean(t)):
                    if c == tcols[0] or not major:
                        major, sub, sub_col = loc_clean(t), "", None
                    else:
                        sub, sub_col = loc_clean(t), c
                    container = ""
                    continue
            *prefixes, (icol, name) = texts
            if len(prefixes) >= 2:
                major, sub, sub_col = loc_clean(prefixes[0][1]), loc_clean(prefixes[1][1]), prefixes[1][0]
                container = ""
            elif len(prefixes) == 1:
                pc, pt = prefixes[0]
                if pc == tcols[0] and not SUBLEVEL.match(pt):
                    major, sub, sub_col = loc_clean(pt), "", None
                    container = ""
                else:
                    sub, sub_col = loc_clean(pt), pc
            elif sub_col is not None and icol <= sub_col:
                sub, sub_col = "", None
            if any(key(name).startswith(key(n)) for n in NOT_ITEMS):
                b.review.append([vid, v["sheet"], f"R{r}", name, "not_item"])
                continue
            loc = " ".join(x for x in (major, sub, container) if x)
            if not loc:
                b.review.append([vid, v["sheet"], f"R{r}", name, "noloc"])
                loc = "場所不明"
            if not has_qty:
                note = "数量記載なし"
                b.review.append([vid, v["sheet"], f"R{r}", name, "noqty"])
            if insp:
                b.review.append([vid, v["sheet"], f"R{r}", f"{name} [{' '.join(insp)}]", "inspection"])
            b.add(vid, loc, name, qty, note)
    for e in v.get("extra", []):
        b.add(vid, loc_clean(e["location"]), clean(e["name"]), e["qty"], "")
    # 比較用シート(旧版など)にだけある品名を照合リストに出す
    known = {key(b.items[k][1]) for k in b.items if any(i == b.items[k][0] and l.startswith(vid) for i, l in b.stock)}
    for sheet in v.get("compare_sheets", []):
        for row in wb[sheet].iter_rows():
            for c in row:
                if not isinstance(c.value, str):
                    continue
                t = clean(c.value)
                if (len(t) <= 2 or HEADING.search(loc_clean(t)) or HEADER_ROW.search(t) or INSPECTION.search(t)
                        or parse_qty([t])[0] or t in ("数量", "異常", "適量")):
                    continue
                if key(t) not in known:
                    b.review.append([vid, sheet, c.coordinate, t, "sheet2_only"])


def pick(rules, text, default):
    for name, pat in rules:
        if re.search(pat, text):
            return name
    return default


def read_csv(path):
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def keep_existing(b, existing):
    """既存の data/ と同じ品名・場所名には同じ番号を振り直し、手入力した欄(位置・写真・通称・分類)を引き継ぐ。"""
    old_items = {key(r["name"]): r for r in read_csv(existing / "items.csv")}
    old_locs = {(r["vehicle_id"], r["name"]): r for r in read_csv(existing / "locations.csv")}
    next_item = max(int(r["item_id"][1:]) for r in old_items.values()) + 1
    item_map, items = {}, {}
    for k, (iid, name) in b.items.items():
        if k in old_items:
            new = old_items[k]["item_id"]
        else:
            new, next_item = f"i{next_item:04d}", next_item + 1
        item_map[iid] = new
        items[k] = (new, name)
    loc_map, locs = {}, {}
    next_loc = {}
    for (vid, _), r in old_locs.items():
        next_loc[vid] = max(next_loc.get(vid, 0), int(r["location_id"].split("-l")[1]))
    renamed = CONFIG.get("renamed_locations", {})
    for (vid, name), lid in b.locations.items():
        before = renamed.get(vid, {}).get(name)  # 点検表で改名された場所は旧名の番号を使う
        if (vid, name) not in old_locs and before and (vid, before) in old_locs:
            old_locs[(vid, name)] = old_locs.pop((vid, before))
        if (vid, name) in old_locs:
            new = old_locs[(vid, name)]["location_id"]
        else:
            next_loc[vid] = next_loc.get(vid, 0) + 1
            new = f"{vid}-l{next_loc[vid]:02d}"
        loc_map[lid] = new
        locs[(vid, name)] = new
    b.items, b.locations = items, locs
    b.stock = {(item_map[i], loc_map[l]): v for (i, l), v in b.stock.items()}
    b.kept_items, b.kept_locs = old_items, old_locs


def write(b, out):
    out.mkdir(parents=True, exist_ok=True)
    kept_items, kept_locs = getattr(b, "kept_items", {}), getattr(b, "kept_locs", {})
    def loc_row(lid, vid, name):
        o = kept_locs.get((vid, name))
        if o:
            return [lid, vid, name, o["group"], o["hotspot"], o["photo"]]
        return [lid, vid, name, pick(CONFIG["location_groups"], name, CONFIG["default_group"]), "", ""]
    def item_row(iid, name):
        o = kept_items.get(key(name))
        if o:
            return [iid, name, o["aliases"], o["category"], o["photo"], o["model"]]
        return [iid, name, reading(name), pick(CONFIG["categories"], name, CONFIG["default_category"]), "", ""]
    def w(name, header, rows):
        with open(out / name, "w", encoding="utf-8", newline="") as f:
            cw = csv.writer(f)
            cw.writerow(header)
            cw.writerows(rows)
    w("vehicles.csv", ["vehicle_id", "name", "type", "model"],
      [[v["id"], v["name"], v["type"], ""] for v in b.vehicles])
    w("locations.csv", ["location_id", "vehicle_id", "name", "group", "hotspot", "photo"],
      [loc_row(lid, vid, name) for (vid, name), lid in b.locations.items()])
    w("items.csv", ["item_id", "name", "aliases", "category", "photo", "model"],
      [item_row(iid, name) for iid, name in b.items.values()])
    w("stock.csv", ["item_id", "location_id", "qty", "qty_note"],
      [[iid, lid, q, n] for (iid, lid), (q, n) in b.stock.items()])
    w("review.csv", ["vehicle_id", "sheet", "cell", "raw", "reason"], b.review)


def main(src, out, existing=None):
    b = Builder()
    for v in CONFIG["vehicles"]:
        read_vehicle(b, v, Path(src))
    if existing:
        keep_existing(b, Path(existing))
    write(b, Path(out))
    print(f"vehicles={len(b.vehicles)} locations={len(b.locations)} items={len(b.items)} "
          f"stock={len(b.stock)} review={len(b.review)}")


if __name__ == "__main__":
    args = sys.argv[1:]
    existing = None
    if "--existing" in args:
        i = args.index("--existing")
        existing = args[i + 1]
        del args[i:i + 2]
    main(args[0], args[1], existing)
